"""
tests.py
Comprehensive unit tests for Talent Align backend including:
- Data Engineering & Excel Automation Architecture (TalentAlignmentExcelReportGenerator)
- Enterprise Stream Normalization & Standardization (10 allowed streams)
- JD Source Tracking (Direct BU vs RMG)
- API endpoints for Excel Export and Transformation
"""

import io
import pandas as pd
import openpyxl
from django.test import TestCase
from django.contrib.auth import get_user_model
from django.urls import reverse
from rest_framework.test import APITestCase
from rest_framework import status

from .models import Job, UserInfo, ProfileRecord, Course, InterviewLock, Match, Strength
from .stream_constants import (
    ALLOWED_STREAMS,
    normalize_to_standard_stream,
    is_valid_stream
)
from .excel_generator import TalentAlignmentExcelReportGenerator

User = get_user_model()


class StreamNormalizationTests(TestCase):
    """Unit tests for the 10 allowed enterprise streams and normalizers."""

    def test_allowed_streams_count(self):
        self.assertEqual(len(ALLOWED_STREAMS), 10)
        expected_streams = [
            'AI Engineering', 'Angular', 'Devops', 'DotNet', 'PLSQL',
            'SpringBoot', 'Test Automation', 'Cloud', 'Cyber Security', 'Data Engineering'
        ]
        self.assertEqual(ALLOWED_STREAMS, expected_streams)

    def test_normalize_synonyms(self):
        # AI Engineering
        self.assertEqual(normalize_to_standard_stream('genai'), 'AI Engineering')
        self.assertEqual(normalize_to_standard_stream('Machine Learning'), 'AI Engineering')
        self.assertEqual(normalize_to_standard_stream('ai'), 'AI Engineering')

        # Angular
        self.assertEqual(normalize_to_standard_stream('angularjs'), 'Angular')
        self.assertEqual(normalize_to_standard_stream('Angular 2+'), 'Angular')

        # Devops
        self.assertEqual(normalize_to_standard_stream('ci/cd'), 'Devops')
        self.assertEqual(normalize_to_standard_stream('kubernetes'), 'Devops')

        # DotNet
        self.assertEqual(normalize_to_standard_stream('c#'), 'DotNet')
        self.assertEqual(normalize_to_standard_stream('.net core'), 'DotNet')
        self.assertEqual(normalize_to_standard_stream('asp.net'), 'DotNet')

        # PLSQL
        self.assertEqual(normalize_to_standard_stream('oracle plsql'), 'PLSQL')
        self.assertEqual(normalize_to_standard_stream('t-sql'), 'PLSQL')

        # SpringBoot
        self.assertEqual(normalize_to_standard_stream('java spring boot'), 'SpringBoot')
        self.assertEqual(normalize_to_standard_stream('spring framework'), 'SpringBoot')

        # Test Automation
        self.assertEqual(normalize_to_standard_stream('selenium'), 'Test Automation')
        self.assertEqual(normalize_to_standard_stream('qa automation'), 'Test Automation')

        # Cloud
        self.assertEqual(normalize_to_standard_stream('aws'), 'Cloud')
        self.assertEqual(normalize_to_standard_stream('azure'), 'Cloud')

        # Cyber Security
        self.assertEqual(normalize_to_standard_stream('infosec'), 'Cyber Security')
        self.assertEqual(normalize_to_standard_stream('penetration testing'), 'Cyber Security')

        # Data Engineering
        self.assertEqual(normalize_to_standard_stream('pyspark'), 'Data Engineering')
        self.assertEqual(normalize_to_standard_stream('etl pipeline'), 'Data Engineering')

    def test_normalize_fallback_others(self):
        self.assertEqual(normalize_to_standard_stream('Unknown Skill XYZ'), 'Others')
        self.assertEqual(normalize_to_standard_stream(''), 'Others')
        self.assertEqual(normalize_to_standard_stream(None), 'Others')


class ExcelGeneratorUnitTests(TestCase):
    """Unit tests for the TalentAlignmentExcelReportGenerator data engineering engine."""

    def setUp(self):
        # Sample raw dataset mimicking business data
        self.sample_records = [
            {
                'BG': 'CBG', 'ISU/HSU': 'Banking', 'Project Name': 'Core Banking System',
                'Stream': 'SpringBoot', 'Location': 'Hyderabad', 'Role': 'Developer',
                'Project SPOC Name': 'Alice Smith', 'Project SPOC Emp ID': 'EMP101',
                'RMG Head': 'Robert Brown', 'RGS ID': 'RGS-001', 'Demand Count': 10,
                'Shared By': 'Direct BU', 'Pending/Confirmed': 'Confirmed', 'Supply Completed': 8,
                'Batch': 'Batch-2026-A'
            },
            {
                'BG': 'Growth Markets', 'ISU/HSU': 'Insurance', 'Project Name': 'Claims Processing',
                'Stream': 'DotNet', 'Location': 'Bangalore', 'Role': 'Developer',
                'Project SPOC Name': 'Bob Jones', 'Project SPOC Emp ID': 'EMP102',
                'RMG Head': 'Robert Brown', 'RGS ID': 'RGS-002', 'Demand Count': 5,
                'Shared By': 'RMG', 'Pending/Confirmed': 'Pending', 'Supply Completed': 3,
                'Batch': 'Batch-2026-A'
            },
            {
                'BG': 'Tech SS', 'ISU/HSU': 'Cloud Ops', 'Project Name': 'Cloud Migration',
                'Stream': 'Cloud', 'Location': 'Pune', 'Role': 'Tech Support',
                'Project SPOC Name': 'Charlie Lee', 'Project SPOC Emp ID': 'EMP103',
                'RMG Head': 'Mary White', 'RGS ID': 'RGS-003', 'Demand Count': 4,
                'Shared By': 'Direct BU', 'Pending/Confirmed': 'Confirmed', 'Supply Completed': 4,
                'Batch': 'Batch-2026-A'
            },
            {
                'BG': 'Japan Ops', 'ISU/HSU': 'Testing COE', 'Project Name': 'Automated Regression',
                'Stream': 'Test Automation', 'Location': 'Chennai', 'Role': 'QA',
                'Project SPOC Name': 'David Green', 'Project SPOC Emp ID': 'EMP104',
                'RMG Head': 'Mary White', 'RGS ID': 'RGS-004', 'Demand Count': 6,
                'Shared By': 'RMG', 'Pending/Confirmed': 'Confirmed', 'Supply Completed': 6,
                'Batch': 'Batch-2026-A'
            }
        ]

    def test_section1_bg_demand_supply(self):
        generator = TalentAlignmentExcelReportGenerator(self.sample_records)
        df_sec1 = generator.compute_section1_bg_demand_supply()
        
        self.assertFalse(df_sec1.empty)
        self.assertIn('BG', df_sec1.columns)
        self.assertIn('Demand Count', df_sec1.columns)
        self.assertIn('Supply Completed', df_sec1.columns)
        
        # Check totals
        total_demand = df_sec1['Demand Count'].sum()
        total_supply = df_sec1['Supply Completed'].sum()
        self.assertEqual(total_demand, 25)
        self.assertEqual(total_supply, 21)

    def test_section2_role_demand_supply(self):
        generator = TalentAlignmentExcelReportGenerator(self.sample_records)
        df_sec2 = generator.compute_section2_role_demand_supply()
        
        self.assertIn('Role', df_sec2.columns)
        dev_row = df_sec2[df_sec2['Role'] == 'Developer']
        self.assertEqual(dev_row['Demand Count'].values[0], 15)
        self.assertEqual(dev_row['Supply Completed'].values[0], 11)

    def test_section3_stream_fulfilment(self):
        generator = TalentAlignmentExcelReportGenerator(self.sample_records)
        df_sec3 = generator.compute_section3_stream_fulfilment()
        
        self.assertIn('Stream', df_sec3.columns)
        self.assertIn('Same Stream Mapping %', df_sec3.columns)
        
        # SpringBoot demand should be 10, same mapped 8
        sb_row = df_sec3[df_sec3['Stream'] == 'SpringBoot']
        self.assertEqual(sb_row['Demand Count'].values[0], 10)
        self.assertEqual(sb_row['Mapped To Same Stream'].values[0], 8)
        self.assertEqual(sb_row['Same Stream Mapping %'].values[0], 0.8)

    def test_section4_cross_stream_matrix(self):
        generator = TalentAlignmentExcelReportGenerator(self.sample_records)
        df_matrix = generator.compute_section4_cross_stream_matrix()
        
        self.assertIn('Total Mapped', df_matrix.columns)
        self.assertIn('SpringBoot', df_matrix.index)
        self.assertEqual(df_matrix.loc['SpringBoot', 'SpringBoot'], 8)

    def test_section5_demand_met(self):
        generator = TalentAlignmentExcelReportGenerator(self.sample_records)
        sec5 = generator.compute_section5_demand_met()
        
        self.assertEqual(sec5['total_demand'], 25)
        self.assertEqual(sec5['total_supply'], 21)
        self.assertEqual(sec5['unfulfilled'], 4)
        self.assertAlmostEqual(sec5['decimal_ratio'], 0.84, places=2)
        self.assertAlmostEqual(sec5['percentage_val'], 84.0, places=1)

    def test_section6_rmg_summary(self):
        generator = TalentAlignmentExcelReportGenerator(self.sample_records)
        sec6 = generator.compute_section6_rmg_summary()
        
        # RMG: RGS-002 (5 demand, 3 supply) + RGS-004 (6 demand, 6 supply) = 11 demand, 9 supply
        self.assertEqual(sec6['rmg_demand'], 11)
        self.assertEqual(sec6['rmg_supply'], 9)
        
        # Direct BU: RGS-001 (10 demand, 8 supply) + RGS-003 (4 demand, 4 supply) = 14 demand, 12 supply
        self.assertEqual(sec6['direct_bu_demand'], 14)
        self.assertEqual(sec6['direct_bu_supply'], 12)

    def test_workbook_structure_and_two_sheets(self):
        generator = TalentAlignmentExcelReportGenerator(self.sample_records)
        wb = generator.generate_workbook()
        
        sheet_names = wb.sheetnames
        self.assertEqual(len(sheet_names), 2)
        self.assertEqual(sheet_names[0], "EXECUTIVE DASHBOARD")
        self.assertEqual(sheet_names[1], "RMG REQUIREMENT")

        # Check Sheet 2 columns
        ws2 = wb["RMG REQUIREMENT"]
        headers = [ws2.cell(row=1, column=col).value for col in range(1, 16)]
        self.assertEqual(headers, [
            "BG", "ISU/HSU", "Project Name", "Stream", "Location", "Role",
            "Project SPOC Name", "Project SPOC Emp ID", "RMG Head", "RGS ID",
            "Demand Count", "Shared By", "Pending/Confirmed", "Supply Completed", "Batch"
        ])
        # Verify 4 data rows + 1 header row = 5 rows
        self.assertEqual(ws2.max_row, 5)

        # Export to BytesIO and verify openpyxl can reload it
        buf = generator.export_to_bytes()
        self.assertGreater(buf.getbuffer().nbytes, 1000)
        reloaded_wb = openpyxl.load_workbook(buf)
        self.assertEqual(reloaded_wb.sheetnames, ["EXECUTIVE DASHBOARD", "RMG REQUIREMENT"])


class APITalentAlignmentTests(APITestCase):
    """Integration tests for Job creation with shared_by and Excel export endpoints."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username='admin_test',
            email='admin@example.com',
            password='Password@123',
            role='admin'
        )
        self.client.force_authenticate(user=self.admin)

        # Create sample job in database
        self.job = Job.objects.create(
            project_name='Payment Gateway Modernization',
            location='Hyderabad',
            demand_id='DMD-991',
            skills='Java, Spring Boot, Microservices',
            stream='SpringBoot',
            openings=5,
            bg='BFSI',
            isu_hsu='Banking',
            role='Developer',
            shared_by='Direct BU',
            rmg_head='Jane RMG',
            spoc_name='John SPOC',
            spoc_emp_id='EMP991',
            rgs_id='RGS-991',
            filled=4,
            status='active',
            batch_name='Batch-2026-Test'
        )

    def test_create_job_with_shared_by_and_standard_stream(self):
        url = reverse('job-list')
        data = {
            'project_name': 'Cloud Infrastructure Security',
            'location': 'Bangalore',
            'demand_id': 'DMD-1002',
            'skills': 'AWS, Terraform, IAM',
            'stream': 'Cloud',
            'openings': 3,
            'bg': 'Tech SS',
            'isu_hsu': 'Security',
            'role': 'Developer',
            'shared_by': 'RMG',
            'spoc_name': 'Alice Manager',
            'spoc_emp_id': 'EMP1002',
            'rmg_head': 'Bob Head',
            'rgs_id': 'RGS-1002',
            'batch_name': 'Batch-2026-Test'
        }
        res = self.client.post(url, data, format='json')
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        created_job = Job.objects.get(rgs_id='RGS-1002')
        self.assertEqual(created_job.shared_by, 'RMG')
        self.assertEqual(created_job.stream, 'Cloud')

    def test_export_talent_alignment_excel_endpoint(self):
        url = reverse('talent-alignment-excel-export')
        res = self.client.get(url)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('talent_alignment_executive_dashboard', res['Content-Disposition'])

        # Load workbook from response content
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertEqual(wb.sheetnames, ["EXECUTIVE DASHBOARD", "RMG REQUIREMENT"])

    def test_parse_jd_extracts_standard_stream_and_shared_by(self):
        url = reverse('parse-jd')
        raw_text = """
        Requirement from RMG: We need 4 Senior Spring Boot Developers for Core Banking Project in Hyderabad.
        Required Skills: Java, Spring Boot, Microservices, REST.
        Role: Developer, BG: BFSI, ISU: Banking, SPOC: Sarah (EMP401).
        """
        res = self.client.post(url, {'text': raw_text}, format='json')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn('stream', res.data)
        self.assertEqual(res.data['stream'], 'SpringBoot')
        self.assertEqual(res.data['shared_by'], 'RMG')